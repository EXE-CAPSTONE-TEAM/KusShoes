"""Reports hub core (RP-1): registry, renderers, escaping, GMT+7 window, generate() guards.
No database: generate() runs against a fake session and a recorded audit call."""

import csv
import importlib
import io
from datetime import UTC, date, datetime, timedelta
from types import SimpleNamespace

import pytest
from openpyxl import load_workbook

from app.config import settings
from app.exceptions import (
    ReportForbidden,
    ReportFormatUnsupported,
    ReportNotFound,
    ReportRangeInvalid,
    ReportRangeTooLarge,
    ReportTooManyRows,
)
from app.services import reports
from app.services.period_service import GMT7
from app.services.reports import registry, render
from app.services.reports.registry import Column, ReportSpec

TRIGGERS = ["=", "+", "-", "@", "\t", "\r"]
DOMAIN_MODULES = ["finance", "billing", "users", "design", "ops", "trust"]


@pytest.fixture
def fresh_registry(monkeypatch):
    monkeypatch.setattr(registry, "_REGISTRY", {})
    return registry


def _spec(report_id="demo", domain="tai-chinh", rows=None, **kwargs) -> ReportSpec:
    async def build(db, ctx):
        return rows if rows is not None else [["a", 1]]

    defaults = dict(
        id=report_id,
        domain=domain,
        title="Báo cáo thử",
        description="d",
        columns=(Column("Tên"), Column("Số", "int")),
        build=build,
    )
    defaults.update(kwargs)
    return ReportSpec(**defaults)


class FakeDb:
    def __init__(self):
        self.commits = 0

    async def commit(self):
        self.commits += 1


@pytest.fixture
def audit_calls(monkeypatch):
    calls = []

    async def fake_record(db, actor, action, **kwargs):
        calls.append((actor, action, kwargs))

    monkeypatch.setattr(reports, "record_audit", fake_record)
    return calls


ADMIN = SimpleNamespace(id="a1", role="admin")
STAFF = SimpleNamespace(id="s1", role="staff")


# --- registry ---------------------------------------------------------------------------


def test_duplicate_id_is_rejected(fresh_registry):
    fresh_registry.register(_spec("x"))
    with pytest.raises(ValueError, match="Duplicate"):
        fresh_registry.register(_spec("x", domain="nguoi-dung"))


def test_unknown_domain_and_bad_kind_rejected():
    with pytest.raises(ValueError):
        _spec(domain="khong-co")
    with pytest.raises(ValueError):
        Column("x", "money")  # type: ignore[arg-type]
    with pytest.raises(ValueError):
        _spec(formats=("pdf",))  # csv + xlsx are mandatory


def test_specs_listed_in_domain_order(fresh_registry):
    fresh_registry.register(_spec("late", domain="kiem-duyet-bao-mat"))
    fresh_registry.register(_spec("first", domain="tai-chinh"))
    fresh_registry.register(_spec("first-b", domain="tai-chinh"))
    assert [s.id for s in fresh_registry.all_specs()] == ["first", "first-b", "late"]


@pytest.mark.parametrize("name", DOMAIN_MODULES)
def test_domain_module_skeleton_exists(name):
    module = importlib.import_module(f"app.services.reports.{name}")
    assert module.DOMAIN in registry.DOMAIN_KEYS


def test_six_domains_declared_once():
    assert len(registry.DOMAIN_KEYS) == len(set(registry.DOMAIN_KEYS)) == 6
    owners = {importlib.import_module(f"app.services.reports.{n}").DOMAIN for n in DOMAIN_MODULES}
    assert owners == set(registry.DOMAIN_KEYS)


def test_catalogue_has_30_unique_ids_and_six_domains_of_at_least_five():
    specs = registry.all_specs()
    assert len({s.id for s in specs}) == len(specs) == 30
    by_domain = {key: [s for s in specs if s.domain == key] for key in registry.DOMAIN_KEYS}
    assert all(len(v) >= 5 for v in by_domain.values())


def test_catalogue_flags_sensitive_as_locked_for_staff(fresh_registry):
    fresh_registry.register(_spec("open"))
    fresh_registry.register(_spec("secret", sensitive=True))
    staff = {e.spec.id: e.locked for e in reports.catalogue("staff")}
    admin = {e.spec.id: e.locked for e in reports.catalogue("admin")}
    assert staff == {"open": False, "secret": True}
    assert admin == {"open": False, "secret": False}


# --- escaping ---------------------------------------------------------------------------


@pytest.mark.parametrize("trigger", TRIGGERS)
def test_safe_cell_prefixes_every_trigger(trigger):
    assert render.safe_cell(f"{trigger}1+1") == f"'{trigger}1+1"


@pytest.mark.parametrize("value", [-5, 0, 12.5, "abc", "", None, "x=1"])
def test_safe_cell_leaves_numbers_and_plain_text(value):
    assert render.safe_cell(value) == value


@pytest.mark.parametrize("trigger", TRIGGERS)
def test_csv_escapes_formula_triggers(trigger):
    raw = render.render_csv(["h"], [[f"{trigger}cmd|' /C calc'!A0"], [-7]])
    parsed = list(csv.reader(io.StringIO(raw.decode("utf-8-sig"))))
    assert parsed[1][0] == f"'{trigger}cmd|' /C calc'!A0"
    assert parsed[2][0] == "-7"  # numbers are not text-escaped


@pytest.mark.parametrize("trigger", TRIGGERS)
def test_xlsx_escaped_cell_is_never_a_formula(trigger):
    data = render.render_xlsx("T", ["h"], [[f"{trigger}SUM(A1)"], [-7]])
    sheet = load_workbook(io.BytesIO(data)).worksheets[0]
    cell = sheet["A2"]
    assert cell.data_type != "f"
    # XML normalises a bare CR to LF on reload; the apostrophe guard is what matters.
    expected = f"'{trigger}SUM(A1)"
    assert cell.value.splitlines() == expected.splitlines()
    assert sheet["A3"].value == -7


def test_unescaped_formula_would_be_a_formula_control():
    """Guard that the xlsx assertion above is meaningful: openpyxl does treat '=' as formula."""
    from openpyxl import Workbook

    sheet = Workbook().active
    sheet["A1"] = "=1+1"
    assert sheet["A1"].data_type == "f"


# --- renderers --------------------------------------------------------------------------


def test_csv_has_utf8_bom_and_vietnamese():
    raw = render.render_csv(["Mã", "Tên khách"], [["KUS-1", "Nguyễn Văn Á"]])
    assert raw.startswith(b"\xef\xbb\xbf")
    assert "Nguyễn Văn Á" in raw.decode("utf-8")


def test_xlsx_sheet_one_is_table_and_info_sheet_second():
    info = [("Tên báo cáo", "T"), ("Số dòng", 2)]
    book = load_workbook(io.BytesIO(render.render_xlsx("Sổ", ["A", "B"], [[1, 2], [3, 4]], info)))
    assert book.sheetnames == ["Sổ", "Thông tin"]
    assert [c.value for c in book.worksheets[0][1]] == ["A", "B"]
    assert [[c.value for c in r] for r in book["Thông tin"].iter_rows()] == [
        ["Tên báo cáo", "T"],
        ["Số dòng", 2],
    ]


def test_xlsx_without_info_has_single_sheet_and_safe_title():
    book = load_workbook(io.BytesIO(render.render_xlsx("a/b:c[d]" + "x" * 40, ["A"], [])))
    assert len(book.sheetnames) == 1
    assert len(book.sheetnames[0]) <= 31
    assert not set("[]:*?/\\") & set(book.sheetnames[0])


def test_pdf_renders():
    assert render.render_pdf("Sổ", ["A"], [["x"]]).startswith(b"%PDF")


def test_format_value_by_kind():
    moment = datetime(2026, 3, 31, 17, 0, tzinfo=UTC)  # = 01/04/2026 00:00 GMT+7
    assert render.format_value(moment, "date") == "01/04/2026"
    assert render.format_value(moment, "datetime") == "01/04/2026 00:00"
    assert render.format_value(date(2026, 4, 1), "date") == "01/04/2026"
    assert render.format_value(0.123, "percent") == "12.3%"
    assert render.format_value(True, "bool") == "Có"
    assert render.format_value(None, "int") == ""
    assert render.format_value(1500, "vnd") == 1500


# --- GMT+7 window -----------------------------------------------------------------------

NOW = datetime(2026, 4, 10, 3, 0, tzinfo=UTC)


def test_window_boundary_at_gmt7_midnight():
    ctx = reports.build_context(_spec(), "admin", date(2026, 3, 1), date(2026, 3, 31), now=NOW)
    inside_last = datetime(2026, 3, 31, 16, 59, 59, tzinfo=UTC)  # 23:59:59 on 31/03 GMT+7
    outside = datetime(2026, 3, 31, 17, 0, tzinfo=UTC)  # 00:00 on 01/04 GMT+7
    assert ctx.start_utc <= inside_last < ctx.end_utc
    assert not ctx.start_utc <= outside < ctx.end_utc
    first = datetime(2026, 2, 28, 17, 0, tzinfo=UTC)  # 00:00 on 01/03 GMT+7
    before = first - timedelta(seconds=1)
    assert ctx.start_utc <= first < ctx.end_utc
    assert not ctx.start_utc <= before < ctx.end_utc


def test_default_range_is_last_30_days_in_business_time():
    # 2026-04-10 20:00 UTC is already 11/04 in GMT+7
    ctx = reports.build_context(
        _spec(), "admin", None, None, now=datetime(2026, 4, 10, 20, 0, tzinfo=UTC)
    )
    assert ctx.end == date(2026, 4, 11)
    assert (ctx.end - ctx.start).days == 29


def test_range_validation():
    spec = _spec(max_range_days=10)
    reports.build_context(spec, "admin", date(2026, 3, 1), date(2026, 3, 10), now=NOW)  # 10 days
    with pytest.raises(ReportRangeTooLarge):
        reports.build_context(spec, "admin", date(2026, 3, 1), date(2026, 3, 11), now=NOW)
    with pytest.raises(ReportRangeInvalid):
        reports.build_context(spec, "admin", date(2026, 3, 5), date(2026, 3, 4), now=NOW)


def test_default_cap_comes_from_settings():
    days = settings.REPORT_MAX_RANGE_DAYS
    end = date(2026, 4, 1)
    reports.build_context(_spec(), "admin", end - timedelta(days=days - 1), end, now=NOW)
    with pytest.raises(ReportRangeTooLarge):
        reports.build_context(_spec(), "admin", end - timedelta(days=days), end, now=NOW)


def test_snapshot_report_ignores_range():
    spec = _spec(date_scoped=False)
    ctx = reports.build_context(spec, "admin", date(2030, 1, 1), date(2020, 1, 1), now=NOW)
    assert ctx.start == ctx.end == date(2026, 4, 10)


# --- generate() -------------------------------------------------------------------------


async def test_generate_unknown_id_is_404(fresh_registry, audit_calls):
    with pytest.raises(ReportNotFound):
        await reports.generate(FakeDb(), ADMIN, "nope", "csv", None, None)
    assert audit_calls == []


async def test_generate_sensitive_blocked_for_staff_before_any_query(fresh_registry, audit_calls):
    async def boom(db, ctx):
        raise AssertionError("build must not run")

    fresh_registry.register(_spec("secret", sensitive=True, build=boom))
    with pytest.raises(ReportForbidden):
        await reports.generate(FakeDb(), STAFF, "secret", "csv", None, None)
    assert audit_calls == []


async def test_generate_pdf_only_when_declared(fresh_registry, audit_calls):
    fresh_registry.register(_spec("plain"))
    fresh_registry.register(_spec("summary", formats=registry.WITH_PDF))
    with pytest.raises(ReportFormatUnsupported):
        await reports.generate(FakeDb(), ADMIN, "plain", "pdf", None, None)
    data, content_type, _ = await reports.generate(FakeDb(), ADMIN, "summary", "pdf", None, None)
    assert data.startswith(b"%PDF") and content_type == "application/pdf"


async def test_generate_row_cap_raises_instead_of_truncating(
    fresh_registry, audit_calls, monkeypatch
):
    monkeypatch.setattr(settings, "REPORT_MAX_ROWS", 3)
    fresh_registry.register(_spec("big", rows=[["r", i] for i in range(4)]))
    fresh_registry.register(_spec("edge", rows=[["r", i] for i in range(3)]))
    with pytest.raises(ReportTooManyRows):
        await reports.generate(FakeDb(), ADMIN, "big", "csv", None, None)
    await reports.generate(FakeDb(), ADMIN, "edge", "csv", None, None)  # exactly at the cap


async def test_generate_audits_commits_and_names_file(fresh_registry, audit_calls):
    fresh_registry.register(_spec("demo", rows=[["=evil", 5], ["ok", 6]]))
    db = FakeDb()
    data, content_type, filename = await reports.generate(
        db, ADMIN, "demo", "xlsx", date(2026, 3, 1), date(2026, 3, 31)
    )
    assert filename == "demo_2026-03-01_2026-03-31.xlsx"
    assert filename.isascii()
    assert content_type == reports.CONTENT_TYPES["xlsx"]
    book = load_workbook(io.BytesIO(data))
    assert book.sheetnames[1] == "Thông tin"
    assert book.worksheets[0]["A2"].value == "'=evil"
    info = {r[0].value: r[1].value for r in book["Thông tin"].iter_rows()}
    assert info["Số dòng"] == 2
    assert "01/03/2026" in info["Kỳ báo cáo (GMT+7)"] and "31/03/2026" in info["Kỳ báo cáo (GMT+7)"]
    assert db.commits == 1
    ((actor, action, kwargs),) = audit_calls
    assert actor is ADMIN and action == "report.export"
    assert kwargs["target_type"] == "report" and kwargs["target_id"] == "demo"
    assert kwargs["payload"] == {
        "format": "xlsx",
        "date_from": "2026-03-01",
        "date_to": "2026-03-31",
        "rows": 2,
    }


async def test_generate_csv_row_count_tracks_input(fresh_registry, audit_calls):
    for n in (0, 1, 7):
        fresh_registry._REGISTRY.pop("r", None)
        fresh_registry.register(_spec("r", rows=[["x", i] for i in range(n)]))
        data, _, _ = await reports.generate(FakeDb(), ADMIN, "r", "csv", None, None)
        assert len(list(csv.reader(io.StringIO(data.decode("utf-8-sig"))))) == n + 1
        assert audit_calls[-1][2]["payload"]["rows"] == n


def test_gmt7_constant_is_plus_seven():
    assert GMT7.utcoffset(None) == timedelta(hours=7)
