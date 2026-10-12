"""Reports hub HTTP contract (RP-2) with dependency_overrides: no database, no real auth."""

import io
from types import SimpleNamespace
from unittest.mock import AsyncMock

import pytest
from httpx import ASGITransport, AsyncClient
from openpyxl import load_workbook

from app.config import settings
from app.dependencies import get_current_admin, get_db
from app.main import app
from app.services import reports
from app.services.reports import registry
from app.services.reports.registry import WITH_PDF, Column, ReportSpec

BASE = "/api/v1/admin/reports"


@pytest.fixture
def audit_calls(monkeypatch):
    calls = []

    async def fake_record(db, actor, action, **kwargs):
        calls.append((actor.role, action, kwargs))

    monkeypatch.setattr(reports, "record_audit", fake_record)
    return calls


@pytest.fixture
def hub(monkeypatch):
    """A small isolated registry; the test chooses the caller's role."""
    monkeypatch.setattr(registry, "_REGISTRY", {})

    async def build(db, ctx):
        return [["=HYPERLINK(1)", 10], ["ok", -3]]

    for report_id, kwargs in {
        "open": {},
        "summary": {"formats": WITH_PDF},
        "secret": {"sensitive": True},
        "snapshot": {"date_scoped": False},
    }.items():
        registry.register(
            ReportSpec(
                id=report_id,
                domain="tai-chinh",
                title=f"Báo cáo {report_id}",
                description="d",
                columns=(Column("Tên"), Column("Số", "int")),
                build=build,
                **kwargs,
            )
        )
    app.dependency_overrides[get_db] = lambda: AsyncMock()

    def as_role(role):
        app.dependency_overrides[get_current_admin] = lambda: SimpleNamespace(id="u1", role=role)

    yield as_role
    app.dependency_overrides.clear()


async def _get(path):
    async with AsyncClient(transport=ASGITransport(app=app), base_url="http://test") as client:
        return await client.get(path)


async def test_catalogue_groups_by_domain_and_flags_locked_for_staff(hub):
    hub("staff")
    body = (await _get(BASE)).json()
    assert [d["key"] for d in body["domains"]] == [k for k, _ in registry.DOMAINS]
    assert body["max_rows"] == settings.REPORT_MAX_ROWS
    items = {r["id"]: r for d in body["domains"] for r in d["reports"]}
    assert items["secret"]["locked"] is True and items["open"]["locked"] is False
    assert items["snapshot"]["max_range_days"] == 0
    assert items["open"]["max_range_days"] == settings.REPORT_MAX_RANGE_DAYS
    assert items["summary"]["formats"] == ["csv", "xlsx", "pdf"]


async def test_catalogue_unlocks_sensitive_for_admin(hub):
    hub("admin")
    items = {r["id"]: r for d in (await _get(BASE)).json()["domains"] for r in d["reports"]}
    assert items["secret"]["locked"] is False


async def test_unknown_id_is_404(hub, audit_calls):
    hub("admin")
    resp = await _get(f"{BASE}/khong-co")
    assert resp.status_code == 404 and resp.json()["code"] == "REPORT_NOT_FOUND"


async def test_staff_forbidden_on_sensitive_admin_allowed(hub, audit_calls):
    hub("staff")
    resp = await _get(f"{BASE}/secret")
    assert resp.status_code == 403 and resp.json()["code"] == "REPORT_FORBIDDEN"
    assert audit_calls == []
    hub("admin")
    assert (await _get(f"{BASE}/secret")).status_code == 200


@pytest.mark.parametrize(
    "query,code",
    [
        ("date_from=2026-03-05&date_to=2026-03-04", "REPORT_RANGE_INVALID"),
        ("date_from=2020-01-01&date_to=2026-03-04", "REPORT_RANGE_TOO_LARGE"),
        ("format=pdf", "REPORT_FORMAT_UNSUPPORTED"),
    ],
)
async def test_bad_requests_are_422_with_code(hub, audit_calls, query, code):
    hub("admin")
    resp = await _get(f"{BASE}/open?{query}")
    assert resp.status_code == 422 and resp.json()["code"] == code
    assert audit_calls == []


async def test_unknown_format_is_framework_422(hub):
    hub("admin")
    assert (await _get(f"{BASE}/open?format=docx")).status_code == 422
    assert (await _get(f"{BASE}/open?date_from=not-a-date")).status_code == 422


async def test_row_cap_is_422_not_truncation(hub, audit_calls, monkeypatch):
    hub("admin")
    monkeypatch.setattr(settings, "REPORT_MAX_ROWS", 1)  # the fixture report yields 2 rows
    resp = await _get(f"{BASE}/open")
    assert resp.status_code == 422 and resp.json()["code"] == "REPORT_TOO_MANY_ROWS"


async def test_csv_download_headers_bom_escaping_and_audit(hub, audit_calls):
    hub("staff")
    resp = await _get(f"{BASE}/open?date_from=2026-03-01&date_to=2026-03-31")
    assert resp.status_code == 200
    assert resp.headers["content-type"].startswith("text/csv")
    assert (
        resp.headers["content-disposition"]
        == 'attachment; filename="open_2026-03-01_2026-03-31.csv"'
    )
    assert resp.content.startswith(b"\xef\xbb\xbf")
    assert "'=HYPERLINK(1)" in resp.content.decode("utf-8-sig")
    ((role, action, kwargs),) = audit_calls
    assert (role, action, kwargs["target_id"]) == ("staff", "report.export", "open")
    assert kwargs["payload"]["rows"] == 2 and kwargs["payload"]["format"] == "csv"


async def test_xlsx_and_pdf_downloads(hub, audit_calls):
    hub("admin")
    xlsx = await _get(f"{BASE}/summary?format=xlsx")
    book = load_workbook(io.BytesIO(xlsx.content))
    assert book.sheetnames[-1] == "Thông tin" and book.worksheets[0]["A2"].data_type != "f"
    pdf = await _get(f"{BASE}/summary?format=pdf")
    assert pdf.content.startswith(b"%PDF")
    assert len(audit_calls) == 2


async def test_snapshot_report_ignores_inverted_range(hub, audit_calls):
    hub("admin")
    resp = await _get(f"{BASE}/snapshot?date_from=2030-01-02&date_to=2030-01-01")
    assert resp.status_code == 200
    assert audit_calls[0][2]["payload"]["date_from"] is None


async def test_unauthenticated_request_is_rejected(monkeypatch):
    monkeypatch.setattr(registry, "_REGISTRY", {})
    app.dependency_overrides.clear()
    app.dependency_overrides[get_db] = lambda: AsyncMock()
    try:
        assert (await _get(BASE)).status_code in (401, 403)
    finally:
        app.dependency_overrides.clear()
