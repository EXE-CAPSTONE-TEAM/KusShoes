import io
import uuid
from datetime import UTC, datetime, timedelta

import bcrypt
import pytest
import pytest_asyncio
from openpyxl import load_workbook

from app.utils.jwt import create_access_token
from tests.test_studio import _create_project, _save


async def _admin(db):
    from app.repositories import user_repo

    admin = await user_repo.create_email_user(
        db,
        email="fb-admin@example.com",
        username="fbadmin",
        password_hash=bcrypt.hashpw(b"Password1", bcrypt.gensalt(rounds=4)).decode(),
        first_name="Admin",
        last_name="Feedback",
    )
    admin.is_verified = True
    admin.role = "admin"
    await db.commit()
    return {"Authorization": f"Bearer {create_access_token(str(admin.id), role='admin')}"}


@pytest_asyncio.fixture
async def studio_user(db, authenticated_user):
    """The default test user, already past their first KusStudio save."""
    authenticated_user.first_studio_save_at = datetime.now(UTC)
    await db.commit()
    return authenticated_user


async def _eligibility(client, headers):
    return (await client.get("/api/v1/feedback/eligibility", headers=headers)).json()


@pytest.mark.asyncio
async def test_feedback_locked_until_first_studio_save(client, auth_headers, service_headers):
    eligibility = await _eligibility(client, auth_headers)
    assert eligibility["has_used_studio"] is False
    assert eligibility["can_submit"] is False

    blocked = await client.post(
        "/api/v1/feedback", headers=auth_headers, json={"rating": 5, "message": "Chưa dùng"}
    )
    assert blocked.status_code == 403
    assert blocked.json()["code"] == "FEEDBACK_REQUIRES_STUDIO_SAVE"

    # Creating a project is not enough; the first save in Studio is what counts.
    project_id = await _create_project(client, auth_headers)
    assert (await _eligibility(client, auth_headers))["has_used_studio"] is False

    assert (await _save(client, service_headers, project_id, {"color": "red"})).status_code == 200
    eligibility = await _eligibility(client, auth_headers)
    assert eligibility["has_used_studio"] is True
    assert eligibility["can_submit"] is True
    created = await client.post(
        "/api/v1/feedback", headers=auth_headers, json={"rating": 5, "message": "Đã dùng"}
    )
    assert created.status_code == 201


@pytest.mark.asyncio
async def test_only_studio_saves_stamp_first_use(client, db, auth_headers, authenticated_user):
    from app.repositories import project_repo

    project = await project_repo.get_by_id(
        db, uuid.UUID(await _create_project(client, auth_headers))
    )

    async def save_as(client_name):
        await project_repo.save_design(
            db,
            project,
            design_config={"saved_by": client_name},
            thumbnail_path=None,
            base_revision=project.current_design_revision,
            author_user_id=authenticated_user.id,
            client=client_name,
        )
        await db.refresh(authenticated_user)
        return authenticated_user.first_studio_save_at

    # Template apply / version restore ("web") and re-crop resets ("desktop") are not a Save.
    assert await save_as("web") is None
    assert await save_as("desktop") is None
    assert await save_as("editor") is not None

    # Later Studio saves never move the first-use moment.
    first_use = datetime(2026, 1, 1, tzinfo=UTC)
    authenticated_user.first_studio_save_at = first_use
    await db.flush()
    assert await save_as("service") == first_use

@pytest.mark.asyncio
async def test_submit_and_cooldown(client, db, auth_headers, studio_user):
    eligibility = await client.get("/api/v1/feedback/eligibility", headers=auth_headers)
    assert eligibility.json()["can_submit"] is True

    created = await client.post(
        "/api/v1/feedback",
        headers=auth_headers,
        json={"rating": 4, "message": "Studio dễ dùng", "marketing_group": "product"},
    )
    assert created.status_code == 201
    assert created.json()["status"] == "new"

    again = await client.post(
        "/api/v1/feedback", headers=auth_headers, json={"rating": 5, "message": "Thêm lần nữa"}
    )
    assert again.status_code == 429
    assert again.json()["code"] == "FEEDBACK_TOO_SOON"
    assert "next_allowed_at" in again.json()

    eligibility = await client.get("/api/v1/feedback/eligibility", headers=auth_headers)
    assert eligibility.json()["can_submit"] is False

    # After 14 days the form opens again.
    from app.repositories import feedback_repo

    latest = await feedback_repo.latest_for_user(db, await _uid(db))
    latest.created_at = datetime.now(UTC) - timedelta(days=15)
    await db.commit()
    ok = await client.post(
        "/api/v1/feedback", headers=auth_headers, json={"rating": 3, "message": "Sau 15 ngày"}
    )
    assert ok.status_code == 201


async def _uid(db):
    from app.repositories import user_repo

    return (await user_repo.get_by_email(db, "profile@example.com")).id


@pytest.mark.asyncio
async def test_rating_bounds_validated(client, auth_headers):
    for rating in (0, 6):
        response = await client.post(
            "/api/v1/feedback", headers=auth_headers, json={"rating": rating, "message": "abc"}
        )
        assert response.status_code == 422


@pytest.mark.asyncio
async def test_admin_triage_summary_and_internal_exclusion(client, db, auth_headers, studio_user):
    from app.repositories import user_repo

    admin_headers = await _admin(db)
    await client.post(
        "/api/v1/feedback",
        headers=auth_headers,
        json={"rating": 2, "message": "Giá hơi cao", "marketing_group": "price"},
    )

    internal = await user_repo.create_email_user(
        db,
        email="staffer@example.com",
        username="staffer",
        password_hash=bcrypt.hashpw(b"Password1", bcrypt.gensalt(rounds=4)).decode(),
        first_name="In",
        last_name="Ternal",
    )
    internal.is_verified = True
    internal.is_internal = True
    internal.first_studio_save_at = datetime.now(UTC)
    await db.commit()
    internal_headers = {"Authorization": f"Bearer {create_access_token(str(internal.id))}"}
    await client.post(
        "/api/v1/feedback", headers=internal_headers, json={"rating": 5, "message": "Nội bộ test"}
    )

    listing = await client.get("/api/v1/admin/feedback", headers=admin_headers)
    assert [f["message"] for f in listing.json()] == ["Giá hơi cao"]
    with_internal = await client.get(
        "/api/v1/admin/feedback?include_internal=true", headers=admin_headers
    )
    assert len(with_internal.json()) == 2

    summary = (await client.get("/api/v1/admin/feedback/summary", headers=admin_headers)).json()
    assert summary["count"] == 1
    assert summary["average_rating"] == 2.0
    assert summary["by_group"] == {"price": 1}
    assert summary["by_status"]["new"] == 1

    feedback_id = listing.json()[0]["id"]
    updated = await client.patch(
        f"/api/v1/admin/feedback/{feedback_id}",
        headers=admin_headers,
        json={"status": "planned", "changed_what": "Thêm gói Basic giá sinh viên"},
    )
    assert updated.status_code == 200
    assert updated.json()["status"] == "planned"

    mine = (await client.get("/api/v1/feedback/mine", headers=auth_headers)).json()
    assert mine[0]["changed_what"] == "Thêm gói Basic giá sinh viên"

    invalid = await client.patch(
        f"/api/v1/admin/feedback/{feedback_id}", headers=admin_headers, json={"status": "bogus"}
    )
    assert invalid.status_code == 422


@pytest.mark.asyncio
async def test_export_xlsx_excludes_internal(client, db, auth_headers, studio_user):
    admin_headers = await _admin(db)
    await client.post(
        "/api/v1/feedback", headers=auth_headers, json={"rating": 5, "message": "Rất tốt"}
    )
    response = await client.get("/api/v1/admin/feedback/export", headers=admin_headers)
    assert response.status_code == 200
    assert "spreadsheetml" in response.headers["content-type"]
    sheet = load_workbook(io.BytesIO(response.content)).active
    rows = list(sheet.iter_rows(values_only=True))
    assert rows[0][0] == "Ngày gửi"
    assert len(rows) == 2
    assert rows[1][4] == "Rất tốt"


@pytest.mark.asyncio
async def test_staff_cannot_update_feedback(client, db, auth_headers, studio_user):
    from app.repositories import user_repo

    staff = await user_repo.create_email_user(
        db,
        email="fb-staff@example.com",
        username="fbstaff",
        password_hash=bcrypt.hashpw(b"Password1", bcrypt.gensalt(rounds=4)).decode(),
        first_name="Staff",
        last_name="Member",
    )
    staff.is_verified = True
    staff.role = "staff"
    await db.commit()
    headers = {"Authorization": f"Bearer {create_access_token(str(staff.id), role='staff')}"}
    created = await client.post(
        "/api/v1/feedback", headers=auth_headers, json={"rating": 4, "message": "OK nhé"}
    )
    response = await client.patch(
        f"/api/v1/admin/feedback/{created.json()['id']}",
        headers=headers,
        json={"status": "reviewed"},
    )
    assert response.status_code == 403
