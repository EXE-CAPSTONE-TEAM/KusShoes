"""SF-16 report generator: CSV / XLSX / PDF on demand.

Each report is a (title, headers, rows) table; the renderers are shared.
Implemented: revenue, users, transactions ("Sổ giao dịch EXE201", BR-106),
channel-funnel (BR-107). Scheduled email delivery is not implemented."""

import asyncio
import csv
import io
from collections import defaultdict
from datetime import UTC, date, datetime, timedelta

from fpdf import FPDF
from openpyxl import Workbook
from sqlalchemy.ext.asyncio import AsyncSession

from app.repositories import analytics_repo
from app.services import analytics_service
from app.services.period_service import GMT7
from app.services.receipt_service import FONT_DIR
from app.utils.http import XLSX_MEDIA_TYPE

_STATUS_LABELS = {
    "paid": "Thành công",
    "failed": "Thất bại",
    "cancelled": "Đã hủy",
    "refunded": "Đã hoàn tiền",
    "awaiting_approval": "Chờ duyệt",
    "pending": "Chờ thanh toán",
}
_METHOD_LABELS = {"payos": "Chuyển khoản VietQR", "momo": "MoMo", "manual": "Thủ công"}
_TIER_LABELS = {"basic": "Basic", "pro": "Pro"}
_CYCLE_LABELS = {"monthly": "tháng", "yearly": "năm"}
_CHANNEL_LABELS = {
    "direct": "Trực tiếp",
    "tiktok": "TikTok",
    "facebook": "Facebook",
    "email": "Email",
    "referral": "Giới thiệu",
    "other": "Khác",
}

Table = tuple[str, list[str], list[list]]


def _bounds(date_from: date | None, date_to: date | None) -> tuple[date, date]:
    end = date_to or analytics_service.business_date(datetime.now(UTC))
    return date_from or end - timedelta(days=29), end


def _to_utc(day: date, *, end_of_day: bool = False) -> datetime:
    moment = datetime(day.year, day.month, day.day, tzinfo=GMT7)
    return moment + timedelta(days=1) if end_of_day else moment


def _channel(value: str | None) -> str:
    return _CHANNEL_LABELS.get(value or "", "Khác")


async def build_report(
    db: AsyncSession, report_type: str, *, date_from: date | None, date_to: date | None
) -> Table:
    start, end = _bounds(date_from, date_to)
    if report_type == "transactions":
        return await _transactions(db, start, end)
    if report_type == "channel-funnel":
        return await _channel_funnel(db, start, end)
    if report_type == "revenue":
        return await _revenue(db, start, end)
    return await _users(db, start, end)


async def _transactions(db: AsyncSession, start: date, end: date) -> Table:
    """BR-106 column order of sổ 07."""
    rows = await analytics_repo.ledger_rows(
        db, start=_to_utc(start), end=_to_utc(end, end_of_day=True)
    )
    payments = analytics_service.build_payments(
        await analytics_repo.paid_invoice_rows(db), await analytics_repo.refund_rows(db)
    )
    repeaters = {
        user_id
        for user_id, history in analytics_service.by_user(payments).items()
        if len({item.paid_at for item in history}) >= 2
    }
    table = []
    for invoice, first, last, account_code, channel in rows:
        snapshot = invoice.receipt_snapshot or {}
        customer = snapshot.get("customer") or account_code
        receipt_file = (
            f"{invoice.receipt_number}-{snapshot.get('file_slug', 'khach')}"
            f"-{snapshot.get('date_code', '')}.pdf"
            if invoice.receipt_number
            else ""
        )
        item = (
            f"Nâng cấp lên {_TIER_LABELS.get(invoice.plan_tier, invoice.plan_tier)}"
            if invoice.is_upgrade
            else f"{_TIER_LABELS.get(invoice.plan_tier, invoice.plan_tier)}"
            f" {_CYCLE_LABELS.get(invoice.billing_cycle, invoice.billing_cycle)}"
        )
        table.append(
            [
                invoice.receipt_number or str(invoice.order_code),
                invoice.created_at.astimezone(GMT7).strftime("%d/%m/%Y"),
                customer,
                _channel(channel),
                item,
                invoice.amount_vnd,
                _STATUS_LABELS.get(invoice.status, invoice.status),
                _METHOD_LABELS.get(invoice.payment_method, invoice.payment_method),
                receipt_file,
                "Có" if invoice.user_id in repeaters else "Chưa",
            ]
        )
    headers = [
        "Mã", "Ngày đặt", "Tên khách", "Kênh đến", "Loại hàng hoá", "Giá tiền",
        "Tình trạng thanh toán", "Hình thức thanh toán", "Tên file biên nhận", "Khách quay lại?",
    ]
    return "Sổ giao dịch EXE201", headers, table


async def _channel_funnel(db: AsyncSession, start: date, end: date) -> Table:
    """BR-107: register → verify email → first saved design → pay, per channel and week."""
    users = await analytics_repo.user_rows(db)
    designers = await analytics_repo.user_ids_with_saved_design(db)
    paying = set(
        analytics_service.by_user(
            analytics_service.build_payments(
                await analytics_repo.paid_invoice_rows(db), await analytics_repo.refund_rows(db)
            )
        )
    )
    cells: dict = defaultdict(lambda: [0, 0, 0, 0])
    for user_id, _email, created_at, verified, channel in users:
        day = analytics_service.business_date(created_at)
        if not start <= day <= end:
            continue
        week = (day - timedelta(days=day.weekday())).isoformat()
        cell = cells[(week, _channel(channel))]
        cell[0] += 1
        cell[1] += bool(verified)
        cell[2] += user_id in designers
        cell[3] += user_id in paying
    table = [
        [week, channel, *counts, f"{counts[3] / counts[1]:.1%}" if counts[1] else ""]
        for (week, channel), counts in sorted(cells.items())
    ]
    headers = [
        "Tuần (bắt đầu)", "Kênh đến", "Đăng ký", "Xác thực email", "Lưu thiết kế đầu",
        "Thanh toán", "Conversion (thanh toán ÷ xác thực)",
    ]
    return "Phễu theo kênh và tuần", headers, table


async def _revenue(db: AsyncSession, start: date, end: date) -> Table:
    data = await analytics_service.get_analytics(db, date_from=start, date_to=end)
    table = [
        ["MRR", data.mrr_vnd],
        ["ARR", data.arr_vnd],
        ["ARPU", data.arpu_vnd],
        ["Khách trả tiền", data.paying_customers],
        ["Doanh thu ghi nhận (kỳ)", data.revenue_vnd.current],
        ["Hoàn tiền (kỳ)", data.refunds_vnd.current],
        ["Khách trả tiền mới (kỳ)", data.new_paying_customers.current],
    ]
    table += [[f"Doanh thu tháng {point.month}", point.revenue_vnd] for point in data.revenue_series]
    table += [[f"Theo gói: {plan.plan_tier}", plan.revenue_vnd] for plan in data.revenue_by_plan]
    return f"Báo cáo doanh thu {start} – {end}", ["Chỉ số", "Giá trị (VND)"], table


async def _users(db: AsyncSession, start: date, end: date) -> Table:
    users = await analytics_repo.user_rows(db)
    table = [
        [
            email,
            created_at.astimezone(GMT7).strftime("%d/%m/%Y"),
            "Có" if verified else "Chưa",
            _channel(channel),
        ]
        for _uid, email, created_at, verified, channel in users
        if start <= analytics_service.business_date(created_at) <= end
    ]
    return (
        f"Báo cáo người dùng {start} – {end}",
        ["Email", "Ngày đăng ký", "Đã xác thực", "Kênh đến"],
        table,
    )


# --- Renderers ------------------------------------------------------------------------


def render_csv(headers: list[str], rows: list[list]) -> bytes:
    buffer = io.StringIO()
    writer = csv.writer(buffer)
    writer.writerow(headers)
    writer.writerows(rows)
    return ("﻿" + buffer.getvalue()).encode("utf-8")  # BOM so Excel reads Vietnamese


def render_xlsx(title: str, headers: list[str], rows: list[list]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = title[:31]
    sheet.append(headers)
    for row in rows:
        sheet.append(row)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def render_pdf(title: str, headers: list[str], rows: list[list]) -> bytes:
    pdf = FPDF(orientation="L", format="A4")
    pdf.add_page()
    pdf.add_font("Roboto", "", str(FONT_DIR / "Roboto-Regular.ttf"))
    pdf.add_font("Roboto", "B", str(FONT_DIR / "Roboto-Bold.ttf"))
    pdf.set_font("Roboto", "B", 13)
    pdf.cell(0, 9, title, new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("Roboto", "", 8)
    with pdf.table(
        first_row_as_headings=True, text_align="LEFT", line_height=5, repeat_headings=1
    ) as table:
        header = table.row()
        for name in headers:
            header.cell(name)
        for values in rows:
            row = table.row()
            for value in values:
                row.cell(str(value))
    return bytes(pdf.output())


CONTENT_TYPES = {
    "csv": "text/csv; charset=utf-8",
    "xlsx": XLSX_MEDIA_TYPE,
    "pdf": "application/pdf",
}


async def generate(
    db: AsyncSession,
    report_type: str,
    fmt: str,
    *,
    date_from: date | None,
    date_to: date | None,
) -> tuple[bytes, str, str]:
    """Returns (content, content_type, filename)."""
    title, headers, rows = await build_report(
        db, report_type, date_from=date_from, date_to=date_to
    )
    renderers = {
        "csv": lambda: render_csv(headers, rows),
        "xlsx": lambda: render_xlsx(title, headers, rows),
        "pdf": lambda: render_pdf(title, headers, rows),
    }
    # fpdf/openpyxl are CPU-bound: keep them off the event loop.
    data = await asyncio.to_thread(renderers[fmt])
    return data, CONTENT_TYPES[fmt], f"{report_type}.{fmt}"
