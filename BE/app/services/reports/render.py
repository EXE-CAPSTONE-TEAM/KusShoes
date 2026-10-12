"""Shared renderers for every report: CSV (UTF-8 BOM), XLSX (write-only, plus a "Thông tin"
sheet) and PDF. Spreadsheet formula injection is neutralised here, once, for all reports."""

import csv
import io
import re
from collections.abc import Sequence
from datetime import date, datetime

from fpdf import FPDF
from openpyxl import Workbook

from app.services.period_service import GMT7
from app.services.receipt_service import FONT_DIR

# OWASP CSV-injection trigger characters (=, +, -, @, TAB, CR); spec.md "Formula-injection".
FORMULA_TRIGGERS = ("=", "+", "-", "@", "\t", "\r")
INFO_SHEET_TITLE = "Thông tin"
CSV_BOM = "﻿"
_SHEET_FORBIDDEN = re.compile(r"[\[\]:*?/\\]")
_DATE_FMT = "%d/%m/%Y"
_DATETIME_FMT = "%d/%m/%Y %H:%M"


def safe_cell(value):
    """Strings starting with a formula trigger get a leading apostrophe; numbers are untouched."""
    if isinstance(value, str) and value.startswith(FORMULA_TRIGGERS):
        return "'" + value
    return value


def format_value(value, kind: str = "text"):
    """Raw Python value -> display value by column kind. Numbers stay numeric (so Excel can sum
    them); dates/datetimes render in GMT+7 as dd/mm/yyyy; None becomes an empty cell."""
    if value is None:
        return ""
    if isinstance(value, bool):
        return "Có" if value else "Không"
    if isinstance(value, datetime):
        moment = value if value.tzinfo else value.replace(tzinfo=GMT7)
        return moment.astimezone(GMT7).strftime(_DATE_FMT if kind == "date" else _DATETIME_FMT)
    if isinstance(value, date):
        return value.strftime(_DATE_FMT)
    if kind == "percent" and isinstance(value, int | float):
        return f"{value:.1%}"
    return value


def format_rows(kinds: Sequence[str], rows: Sequence[Sequence]) -> list[list]:
    return [
        [format_value(value, kinds[i] if i < len(kinds) else "text") for i, value in enumerate(row)]
        for row in rows
    ]


def _sheet_title(title: str) -> str:
    cleaned = _SHEET_FORBIDDEN.sub(" ", title).strip()[:31]
    return cleaned or "Báo cáo"


def render_csv(headers: Sequence[str], rows: Sequence[Sequence]) -> bytes:
    buffer = io.StringIO()
    writer = csv.writer(buffer)
    writer.writerow(headers)
    writer.writerows([safe_cell(value) for value in row] for row in rows)
    return (CSV_BOM + buffer.getvalue()).encode("utf-8")  # BOM so Excel reads Vietnamese


def render_xlsx(
    title: str,
    headers: Sequence[str],
    rows: Sequence[Sequence],
    info: Sequence[tuple[str, object]] | None = None,
) -> bytes:
    """Sheet 1 = header + table. `info` adds sheet 2 "Thông tin" (CR-12 extraction metadata)."""
    workbook = Workbook(write_only=True)
    sheet = workbook.create_sheet(_sheet_title(title))
    sheet.append(list(headers))
    for row in rows:
        sheet.append([safe_cell(value) for value in row])
    if info:
        meta = workbook.create_sheet(INFO_SHEET_TITLE)
        for label, value in info:
            meta.append([safe_cell(label), safe_cell(value)])
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def render_pdf(title: str, headers: Sequence[str], rows: Sequence[Sequence]) -> bytes:
    pdf = FPDF(orientation="L", format="A4")
    pdf.add_page()
    pdf.add_font("Roboto", "", str(FONT_DIR / "Roboto-Regular.ttf"))
    pdf.add_font("Roboto", "B", str(FONT_DIR / "Roboto-Bold.ttf"))
    pdf.set_font("Roboto", "B", 13)
    pdf.cell(0, 9, title, new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("Roboto", "", 8)
    with pdf.table(
        first_row_as_headings=True, text_align="LEFT", line_height=5, repeat_headings=1
    ) as table:
        header = table.row()
        for name in headers:
            header.cell(name)
        for values in rows:
            row = table.row()
            for value in values:
                row.cell(str(value))
    return bytes(pdf.output())
